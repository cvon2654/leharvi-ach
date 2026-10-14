"""Run every solver over a workbook and save the solved copy."""

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl.styles import Alignment

from . import accounts
from .inventory import fill_inventory
from .journal import (book_from_balances, book_from_journal, fill_missing, find_tables)
from .schedules import Inputs, fill_bonds, fill_depreciation, fill_loan
from .statements import (create_balance_sheet, create_equity_statement, create_income_statement,
                         create_ledger, create_trial_balance, existing_statements, guess_company,
                         guess_period_end)
from .totals import fill_templates
from .workbook import (BOLD, FILL_COLOR, TITLE, Context, cell_ref, is_label, norm, set_widths,
                       style_header)

NOTES_SHEET = "Solver Notes"


@dataclass
class Step:
    kind: str            # "fill" or "sheet"
    sheet: str
    cell: str = ""
    note: str = ""
    content: object = None


@dataclass
class Result:
    output: str
    steps: list = field(default_factory=list)
    notes: list = field(default_factory=list)

    @property
    def filled(self):
        return [s for s in self.steps if s.kind == "fill"]

    def summary(self):
        sheets = [s.sheet for s in self.steps if s.kind == "sheet"]
        lines = [f"Filled {len(self.filled)} cell(s)" +
                 (f" and added {len(sheets)} sheet(s): {', '.join(sheets)}" if sheets else ""),
                 f"Saved: {self.output}"]
        for level, msg in self.notes:
            lines.append(f"[{level}] {msg}")
        return "\n".join(lines)


def load_chart_override(ctx):
    """A sheet titled 'Chart of Accounts' with Account | Type columns overrides guesses."""
    for ws in ctx.wb.worksheets:
        if "chart" not in norm(ws.title) and not any(
                is_label(ctx.value(ws, r, c)) and "chart of accounts" in norm(ctx.value(ws, r, c))
                for r in range(1, 4) for c in range(1, 4)):
            continue
        for r in range(1, 15):
            heads = {c: norm(ctx.value(ws, r, c)) for c in range(1, 10)
                     if is_label(ctx.value(ws, r, c))}
            acol = next((c for c, t in heads.items() if t.startswith("account") and
                         "type" not in t), None)
            tcol = next((c for c, t in heads.items() if re.search(r"type|class|category", t)),
                        None)
            if acol and tcol:
                for rr in range(r + 1, ws.max_row + 1):
                    name, typ = ctx.value(ws, rr, acol), ctx.value(ws, rr, tcol)
                    if is_label(name) and is_label(typ):
                        t = accounts.type_from_words(typ)
                        if t:
                            ctx.type_override[accounts.account_key(name)] = t
                ctx.claim(ws, range(1, ws.max_row + 1))
                ctx.info(f"Using account types from '{ws.title}' "
                         f"({len(ctx.type_override)} accounts).")
                return


def _fill_tb_from_journal(ctx, t, book):
    """A blank trial balance template + a journal: carry each account's balance over."""
    ws, p = t.ws, t.pairs[0]
    for ln in t.account_lines:
        if not (ctx.is_empty_cell(ws, ln.row, p.debit) and ctx.is_empty_cell(ws, ln.row,
                                                                             p.credit)):
            continue
        info = book.get(ln.name)
        if info is None:
            ctx.warn(f"{ws.title}: '{ln.name}' is not in the journal - left blank.")
            continue
        net = info.dr_minus_cr()
        ctx.fill(ws, ln.row, p.debit, f"=MAX(0,{net})", f"{ln.name}: debit balance from journal")
        ctx.fill(ws, ln.row, p.credit, f"=MAX(0,-{net})", f"{ln.name}: credit balance from "
                                                          f"journal")


def solve(path, output=None, highlight=True):
    ctx = Context(path, highlight=highlight)
    ctx.template_fills = set()
    ctx.events = []
    _wrap_events(ctx)
    if NOTES_SHEET in ctx.wb.sheetnames:
        del ctx.wb[NOTES_SHEET]
    load_chart_override(ctx)

    tables = []
    for ws in list(ctx.wb.worksheets):
        tables += find_tables(ctx, ws)
    journals = [t for t in tables if t.kind == "journal"]
    tbs = [t for t in tables if t.kind == "tb"]
    sheets10 = [t for t in tables if t.kind == "worksheet"]
    jbook = book_from_journal(ctx, journals) if journals else None

    for t in tables:
        if t.kind == "tb" and jbook is not None:
            _fill_tb_from_journal(ctx, t, jbook)
        fill_missing(ctx, t)

    company = guess_company(ctx)
    period = guess_period_end(ctx, journals)
    book = None
    if sheets10:
        t = sheets10[0]
        pair = t.roles.get("atb") or t.roles.get("tb")
        if pair:
            book = book_from_balances(ctx, t, pair)
    elif tbs:
        adjusted = [t for t in tbs if "adjusted" in t.title_text and "unadjusted" not in
                    t.title_text]
        t = (adjusted or tbs)[-1]
        book = book_from_balances(ctx, t, t.pairs[0])
    elif jbook is not None and len(jbook):
        book = jbook
        create_trial_balance(ctx, book, journals, company, period)
        create_ledger(ctx, book, journals)
    ctx.book = book

    inp = Inputs(ctx)
    ctx.inputs = inp
    fill_depreciation(ctx, inp)
    fill_loan(ctx, inp)
    fill_bonds(ctx, inp)
    fill_inventory(ctx, inp)

    have = existing_statements(ctx) if book else set()
    fill_templates(ctx, final=False)
    if book is not None and len(book):
        if "is" not in have:
            create_income_statement(ctx, book, company, period)
        if "equity" not in have:
            create_equity_statement(ctx, book, company, period)
        if "bs" not in have:
            create_balance_sheet(ctx, book, company, period)
    fill_templates(ctx, final=True)

    steps = _steps(ctx)
    if not steps:
        ctx.warn("Solver did not find anything it knows how to fill. It looks for journals, "
                 "trial balances, worksheets, statements with blank totals, depreciation / "
                 "loan / bond / inventory data, and ratios. Make sure each number has a label "
                 "next to it and the answer cells are empty or contain '?'.")
    _write_notes(ctx, steps)
    out = Path(output) if output else _default_output(path, ctx.keep_vba)
    ctx.wb.calculation.fullCalcOnLoad = True
    ctx.wb.save(out)
    return Result(str(out), steps, ctx.notes)


def _wrap_events(ctx):
    orig_fill, orig_new = ctx.fill, ctx.new_sheet

    def fill(ws, row, col, content, description, fmt=None):
        cell = orig_fill(ws, row, col, content, description, fmt)
        ctx.events.append(Step("fill", ws.title, cell.coordinate, description, content))
        return cell

    def new_sheet(title):
        ws = orig_new(title)
        ctx.events.append(Step("sheet", ws.title, note=f"Solver added the sheet '{ws.title}'"))
        return ws

    ctx.fill, ctx.new_sheet = fill, new_sheet


def _steps(ctx):
    """One step per answer cell (latest content, first position), in solving order."""
    latest = {}
    for e in ctx.events:
        if e.kind == "fill":
            latest[(e.sheet, e.cell)] = e
    out, seen = [], set()
    for e in ctx.events:
        key = (e.sheet, e.cell) if e.kind == "fill" else ("sheet", e.sheet)
        if key in seen:
            continue
        seen.add(key)
        if e.kind == "fill":
            out.append(latest.get(key, e))
        else:
            e.note = f"Solver added the sheet '{e.sheet}' (labels first, then each number)"
            out.append(e)
            out.extend(_sheet_boxes(ctx.wb[e.sheet]))
    return out


def _sheet_boxes(ws):
    """Every formula on a sheet Solver created, as its own step, top to bottom."""
    steps = []
    for row in ws.iter_rows():
        label = next((c.value.strip() for c in row
                      if isinstance(c.value, str) and not c.value.startswith("=")), "")
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                head = ""
                for r in range(c.row - 1, 0, -1):
                    v = ws.cell(row=r, column=c.column).value
                    if isinstance(v, str) and not v.startswith("="):
                        head = v
                        break
                what = " - ".join(x for x in (label, head) if x) or c.coordinate
                steps.append(Step("fill", ws.title, c.coordinate, what, c.value))
    return steps


def _default_output(path, keep_vba):
    p = Path(path)
    return p.with_name(f"{p.stem} - Solved{'.xlsm' if keep_vba else '.xlsx'}")


def _write_notes(ctx, steps):
    ws = ctx.wb.create_sheet(NOTES_SHEET)
    ws.cell(row=1, column=1, value="Solver Notes").font = TITLE
    ws.cell(row=2, column=1, value="Green cells were filled in by Solver. Each answer is a "
                                   "formula, so click a cell to see the work.")
    c = ws.cell(row=2, column=6, value="")
    c.fill = FILL_COLOR
    r = 4
    for col, text in enumerate(["Step", "Sheet", "Cell", "What it is", "Formula / value"], 1):
        ws.cell(row=r, column=col, value=text)
    style_header(ws, r, range(1, 6))
    r += 1
    for i, s in enumerate(steps, 1):
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=s.sheet)
        ws.cell(row=r, column=3, value=s.cell or "(new sheet)")
        ws.cell(row=r, column=4, value=s.note)
        if s.kind == "fill":
            cell = ws.cell(row=r, column=5, value=str(s.content))
            cell.data_type = "s"    # show the formula as text, not a second calculation
        r += 1
    if ctx.notes:
        r += 1
        ws.cell(row=r, column=1, value="Notes and warnings").font = BOLD
        r += 1
        for level, msg in ctx.notes:
            ws.cell(row=r, column=1, value=level.upper())
            m = ws.cell(row=r, column=2, value=msg)
            m.alignment = Alignment(wrap_text=False)
            r += 1
    set_widths(ws, [8, 22, 14, 70, 60])
