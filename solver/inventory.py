"""Inventory costing (periodic): FIFO, LIFO and weighted average."""

import re

from openpyxl.utils import get_column_letter

from .workbook import (BOLD, cell_ref, is_blank, is_label, is_number, norm, set_widths,
                       style_header, write_title)

METHODS = [("fifo", r"\bfifo\b|first[- ]in"), ("lifo", r"\blifo\b|last[- ]in"),
           ("avg", r"average")]
NAMES = {"fifo": "FIFO", "lifo": "LIFO", "avg": "Weighted average"}


def _method(text):
    for key, pat in METHODS:
        if re.search(pat, text):
            return key
    return None


def _find_table(ctx):
    for ws in ctx.wb.worksheets:
        if ws.title in ctx.created:
            continue
        for r in range(1, min(ws.max_row, 300) + 1):
            if ctx.is_claimed(ws, r):
                continue
            roles = {}
            for c in range(1, min(ws.max_column, 15) + 1):
                v = ctx.value(ws, r, c)
                if not isinstance(v, str):
                    continue
                t = norm(v)
                if re.search(r"unit cost|cost per unit|per unit|unit price|price per unit|"
                             r"cost/unit|cost each", t):
                    roles.setdefault("cost", c)
                elif re.search(r"^(units|quantity|qty|no\.? of units|number of units)", t):
                    roles.setdefault("units", c)
                elif re.search(r"total|^cost$|amount", t):
                    roles.setdefault("total", c)
                elif "desc" in t or "item" in t:
                    roles["desc"] = c
                elif re.search(r"date|transaction|^$", t) or "desc" not in roles:
                    roles.setdefault("desc", c)
            if "cost" in roles and "units" in roles:
                return ws, r, roles
    return None


def fill_inventory(ctx, inp):
    found = _find_table(ctx)
    if not found:
        return
    ws, hdr, roles = found
    layers, sales, total_row = [], [], None
    r = hdr + 1
    while r <= hdr + 100:
        desc = ctx.value(ws, r, roles["desc"]) if "desc" in roles else None
        units = ctx.number(ws, r, roles["units"])
        if units is None and not is_label(desc):
            if layers:
                break
            r += 1
            continue
        t = norm(desc) if isinstance(desc, str) else ""
        if re.search(r"total|goods available|available for sale", t):
            total_row = r
            break
        if re.search(r"\bsale|\bsold|\bsales\b|issued|used", t):
            sales.append(r)
        elif units is not None:
            layers.append(r)
        r += 1
    if not layers:
        return
    ctx.claim(ws, range(hdr, (total_row or layers[-1]) + 1))
    U, C = get_column_letter(roles["units"]), get_column_letter(roles["cost"])
    if "total" in roles:
        T = get_column_letter(roles["total"])
        for lr in layers:
            if ctx.is_empty_cell(ws, lr, roles["total"]):
                ctx.fill(ws, lr, roles["total"], f"={U}{lr}*{C}{lr}", "Units x unit cost")
        if total_row:
            if ctx.is_empty_cell(ws, total_row, roles["units"]):
                ctx.fill(ws, total_row, roles["units"], "=" + "+".join(f"{U}{x}" for x in layers),
                         "Units available for sale", "#,##0")
            if ctx.is_empty_cell(ws, total_row, roles["total"]):
                ctx.fill(ws, total_row, roles["total"], "=" + "+".join(f"{T}{x}" for x in layers),
                         "Cost of goods available for sale")

    sold_in = inp.find(r"units sold|sold units|number of units sold|sales in units|"
                       r"units (were )?sold|quantity sold")
    end_in = inp.find(r"(ending|closing) inventory.*units|units.*(ending|on hand|in ending|"
                      r"remaining)|units on hand|physical count")
    avail = "(" + "+".join(cell_ref(ws, x, roles["units"]) for x in layers) + ")"
    if sold_in:
        sold = sold_in.ref()
    elif sales:
        sold = "(" + "+".join(f"ABS({cell_ref(ws, x, roles['units'])})" for x in sales) + ")"
    elif end_in:
        sold = f"({avail}-{end_in.ref()})"
    else:
        ctx.warn("Inventory table found but no units sold / ending units - FIFO/LIFO not computed.")
        return

    out = ctx.new_sheet("Inventory Costing")
    row = write_title(out, ["Inventory Costing (periodic)",
                            f"Layers from '{ws.title}' rows {layers[0]}-{layers[-1]}"])
    heads = ["Layer", "Units", "Unit cost", "Total cost", "FIFO units sold", "FIFO cost sold",
             "LIFO units sold", "LIFO cost sold"]
    for c, h in enumerate(heads, 1):
        out.cell(row=row, column=c, value=h)
    style_header(out, row, range(1, 9))
    first = row + 1
    n = len(layers)
    last = first + n - 1
    money = '#,##0.00_);(#,##0.00)'
    for i, lr in enumerate(layers):
        rr = first + i
        d = ctx.value(ws, lr, roles["desc"]) if "desc" in roles else None
        out.cell(row=rr, column=1, value=str(d) if d is not None else f"Layer {i + 1}")
        out.cell(row=rr, column=2, value="=" + cell_ref(ws, lr, roles["units"]))
        out.cell(row=rr, column=3, value="=" + cell_ref(ws, lr, roles["cost"])).number_format = money
        out.cell(row=rr, column=4, value=f"=B{rr}*C{rr}").number_format = money
        before = f"SUM(B${first}:B{rr - 1})" if i else "0"
        after = f"SUM(B{rr + 1}:B${last})" if rr < last else "0"
        out.cell(row=rr, column=5, value=f"=MAX(0,MIN(B{rr},$B${last + 4}-{before}))")
        out.cell(row=rr, column=6, value=f"=E{rr}*C{rr}").number_format = money
        out.cell(row=rr, column=7, value=f"=MAX(0,MIN(B{rr},$B${last + 4}-{after}))")
        out.cell(row=rr, column=8, value=f"=G{rr}*C{rr}").number_format = money
    s = last + 2
    lines = [
        ("Units available for sale", f"=SUM(B{first}:B{last})", "#,##0"),
        ("Cost of goods available for sale", f"=SUM(D{first}:D{last})", money),
        ("Units sold", f"={sold}", "#,##0"),
        ("Units in ending inventory", f"=B{s}-B{s + 2}", "#,##0"),
        ("FIFO cost of goods sold", f"=SUM(F{first}:F{last})", money),
        ("FIFO ending inventory", f"=B{s + 1}-B{s + 4}", money),
        ("LIFO cost of goods sold", f"=SUM(H{first}:H{last})", money),
        ("LIFO ending inventory", f"=B{s + 1}-B{s + 6}", money),
        ("Weighted-average unit cost", f"=B{s + 1}/B{s}", '#,##0.0000'),
        ("Weighted-average cost of goods sold", f"=B{s + 2}*B{s + 8}", money),
        ("Weighted-average ending inventory", f"=B{s + 3}*B{s + 8}", money),
    ]
    keys = ["units available", "cost of goods available for sale", "units sold", "ending units",
            "fifo cogs", "fifo ending", "lifo cogs", "lifo ending", "avg unit cost", "avg cogs",
            "avg ending"]
    refs = {}
    for i, (label, f, fmt) in enumerate(lines):
        rr = s + i
        out.cell(row=rr, column=1, value=label).font = BOLD
        out.cell(row=rr, column=2, value=f).number_format = fmt
        refs[keys[i]] = cell_ref(out, rr, 2)
    set_widths(out, [34, 14, 12, 14, 16, 16, 16, 16])
    ctx.info(f"Created sheet '{out.title}' (FIFO, LIFO and weighted average).")
    _fill_targets(ctx, refs)


def _fill_targets(ctx, refs):
    """Blank cells labelled e.g. 'FIFO ending inventory', or under a FIFO/LIFO column header."""
    for ws in ctx.wb.worksheets:
        if ws.title in ctx.created:
            continue
        for r in range(1, min(ws.max_row, 300) + 1):
            if ctx.is_claimed(ws, r):
                continue
            label, lc = None, 0
            for c in range(1, min(ws.max_column, 10) + 1):
                v = ctx.value(ws, r, c)
                if is_label(v):
                    label, lc = norm(v), c
                    break
            if not label or len(label) > 60:
                continue
            if re.search(r"ending inventory|inventory, end|closing inventory", label):
                what = "ending"
            elif re.search(r"cost of goods sold|\bcogs\b|cost of sales", label):
                what = "cogs"
            elif re.search(r"average (unit )?cost", label):
                what = "avg unit cost"
            elif re.search(r"goods available", label):
                what = "cost of goods available for sale"
            else:
                continue
            for c in range(lc + 1, min(ws.max_column, lc + 6) + 1):
                if not ctx.is_empty_cell(ws, r, c):
                    continue
                m = _method(label)
                if not m:
                    for rr in range(r - 1, max(0, r - 12), -1):
                        v = ctx.value(ws, rr, c)
                        if isinstance(v, str) and _method(norm(v)):
                            m = _method(norm(v))
                            break
                key = what if what in ("avg unit cost", "cost of goods available for sale") \
                    else (f"{m} {what}" if m else None)
                if key and key in refs:
                    ctx.fill(ws, r, c, "=" + refs[key],
                             f"{NAMES.get(m, '')} {label} (see Inventory Costing sheet)".strip())
                    ctx.claim(ws, [r])
                    if _method(label):
                        break
