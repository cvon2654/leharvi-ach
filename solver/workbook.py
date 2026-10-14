"""Shared helpers: reading cells, writing answers, and the solve context."""

import datetime as _dt
import re

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import accounts

ACCOUNTING_FMT = '#,##0.00_);(#,##0.00);"-"??_)'
RATIO_FMT = "0.00"
PERCENT_FMT = "0.00%"
FILL_COLOR = PatternFill("solid", fgColor="E2F0D9")      # light green = filled by Solver
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=13)
NOTE_FONT = Font(size=9, italic=True)
THIN = Side(style="thin")
DOUBLE = Side(style="double")

PLACEHOLDER_RE = re.compile(
    r"^\s*(\$?\s*\?+|_+|\$\s*_*|\$|x{2,}|-{3,}|\.{3,}|tbd|n/?a\?|\(\s*\?\s*\)|\[\s*\]|\?\?+)\s*$",
    re.I)


class _FormulaNoCache:
    """A formula cell whose computed value was not saved in the file."""

    def __repr__(self):
        return "<formula>"


FORMULA = _FormulaNoCache()


def norm(text):
    """Lower-case a label and strip punctuation so labels compare loosely."""
    if text is None:
        return ""
    t = str(text).lower().replace("’", "'").replace("–", "-").replace("—", "-")
    t = re.sub(r"[$*:;\[\]]", " ", t)
    t = re.sub(r"\s+", " ", t).strip(" .,-")
    return t


def is_placeholder(v):
    return isinstance(v, str) and bool(PLACEHOLDER_RE.match(v))


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def is_blank(v):
    return v is None or (isinstance(v, str) and (v.strip() == "" or is_placeholder(v)))


def is_label(v):
    if not isinstance(v, str) or is_blank(v) or v.startswith("="):
        return False
    return bool(re.search(r"[A-Za-z]", v))


def quote_sheet(title):
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", title):
        return title
    return "'" + title.replace("'", "''") + "'"


def cell_ref(ws, row, col, from_ws=None, absolute=False):
    letter = get_column_letter(col)
    coord = f"${letter}${row}" if absolute else f"{letter}{row}"
    if from_ws is not None and from_ws.title == ws.title:
        return coord
    return f"{quote_sheet(ws.title)}!{coord}"


def number_to_text(n):
    if n is None:
        return ""
    if abs(n - round(n)) < 1e-9:
        return f"{int(round(n)):,}"
    return f"{n:,.2f}"


class Context:
    """Everything one solve run knows: the workbook, answers written, notes."""

    def __init__(self, path, highlight=True):
        self.path = str(path)
        self.keep_vba = self.path.lower().endswith(".xlsm")
        self.wb = load_workbook(self.path, keep_vba=self.keep_vba)
        self.vwb = load_workbook(self.path, data_only=True)
        self.highlight = highlight
        self.filled = []            # (sheet title, coordinate, description, content)
        self.notes = []             # (level, message)
        self.created = []           # titles of sheets Solver added
        self.claimed = {}           # sheet title -> set of rows other solvers own
        self.registry = {}          # concept key -> formula expression (no "=")
        self.registry_src = {}      # concept key -> sheet title it came from
        self.book = None            # accounts ledger (journal.AccountBook)
        self.type_override = {}     # account key -> type from a chart of accounts sheet
        self.inputs = None          # key-value inputs found in the workbook

    # ----- reading -------------------------------------------------------
    def raw(self, ws, row, col):
        return ws.cell(row=row, column=col).value

    def value(self, ws, row, col):
        """Cell value; formulas resolve to their cached result (or FORMULA)."""
        v = ws.cell(row=row, column=col).value
        if isinstance(v, str) and v.startswith("="):
            if ws.title in self.vwb.sheetnames:
                cached = self.vwb[ws.title].cell(row=row, column=col).value
                if cached is not None:
                    return cached
            return FORMULA
        return v

    def number(self, ws, row, col):
        v = self.value(ws, row, col)
        if is_number(v):
            return float(v)
        if isinstance(v, str):
            s = v.strip().replace(",", "").replace("$", "")
            neg = s.startswith("(") and s.endswith(")")
            s = s.strip("()")
            pct = s.endswith("%")
            s = s.rstrip("%")
            try:
                n = float(s)
            except ValueError:
                return None
            n = -n if neg else n
            return n / 100 if pct else n
        return None

    def is_empty_cell(self, ws, row, col):
        return is_blank(ws.cell(row=row, column=col).value)

    def classify(self, name):
        key = accounts.account_key(name)
        if key in self.type_override:
            return self.type_override[key]
        return accounts.classify(name)

    def claim(self, ws, rows):
        self.claimed.setdefault(ws.title, set()).update(rows)

    def is_claimed(self, ws, row):
        return row in self.claimed.get(ws.title, ())

    # ----- writing -------------------------------------------------------
    def fill(self, ws, row, col, content, description, fmt=None):
        cell = ws.cell(row=row, column=col)
        if isinstance(cell, MergedCell):
            for rng in ws.merged_cells.ranges:
                if cell.coordinate in rng:
                    cell = ws.cell(row=rng.min_row, column=rng.min_col)
                    break
        cell.value = content
        if fmt:
            cell.number_format = fmt
        elif cell.number_format == "General" and (
                is_number(content) or (isinstance(content, str) and content.startswith("="))):
            cell.number_format = ACCOUNTING_FMT
        if self.highlight:
            cell.fill = FILL_COLOR
        self.filled.append((ws.title, cell.coordinate, description, content))
        return cell

    def info(self, msg):
        self.notes.append(("info", msg))

    def warn(self, msg):
        self.notes.append(("warning", msg))

    def register(self, key, expr, ws=None, overwrite=False):
        if not key or not expr:
            return
        if overwrite or key not in self.registry:
            self.registry[key] = expr
            self.registry_src[key] = ws.title if ws is not None else None

    def lookup(self, *keys):
        for k in keys:
            if k in self.registry:
                return self.registry[k]
        return None

    def new_sheet(self, title):
        base = title
        n = 2
        while title in self.wb.sheetnames:
            title = f"{base} ({n})"
            n += 1
        ws = self.wb.create_sheet(title)
        self.created.append(title)
        return ws


# ----- helpers for sheets Solver writes ----------------------------------

def write_title(ws, lines, width_cols=3):
    row = 1
    for i, text in enumerate(lines):
        if not text:
            continue
        c = ws.cell(row=row, column=1, value=text)
        c.font = TITLE if i == 0 else BOLD
        c.alignment = Alignment(horizontal="left")
        row += 1
    return row + 1


def style_header(ws, row, cols):
    for col in cols:
        c = ws.cell(row=row, column=col)
        c.font = BOLD
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def underline(ws, row, col, double=False):
    c = ws.cell(row=row, column=col)
    c.border = Border(top=THIN, bottom=DOUBLE if double else None)


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sum_formula(refs):
    """=SUM(...) that stays under Excel's 255-argument limit."""
    refs = list(refs)
    if not refs:
        return "0"
    chunks = [refs[i:i + 200] for i in range(0, len(refs), 200)]
    return "+".join("SUM(" + ",".join(chunk) + ")" for chunk in chunks)


def find_dates(values):
    return [v for v in values if isinstance(v, (_dt.date, _dt.datetime))]
