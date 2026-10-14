"""Find account tables (journals, trial balances, 10-column worksheets),
fill their missing amounts, and collect account balances for the statements."""

import re
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter

from . import accounts
from .workbook import (FORMULA, cell_ref, is_blank, is_label, is_number, norm, sum_formula)

TITLE_STOP = re.compile(r"^(income statement|balance sheet|statement of|trial balance|"
                        r"general journal|journal|general ledger|ledger|worksheet|work sheet)")
EXPLANATION = re.compile(r"^(\(|to record|record|being|explanation|purchased|paid|received|"
                         r"sold|issued|provided|performed|billed|invested|collected|"
                         r"recorded|owner|closing entr|adjusting entr|closed)")


def _is_account_header(t):
    return t.startswith("account") or t in {"description", "particulars", "ledger account",
                                            "title", "account title", "accounts"}


def _is_debit_header(t):
    return t in {"debit", "debits", "dr", "dr."} or t.startswith("debit") or t.startswith("dr ")


def _is_credit_header(t):
    return t in {"credit", "credits", "cr", "cr."} or t.startswith("credit") or \
        t.startswith("cr ")


def _is_date_header(t):
    return t in {"date", "dates", "trans date", "transaction date"} or t.startswith("date")


def _is_ref_header(t):
    return t in {"ref", "pr", "p.r", "post ref", "post. ref", "lf", "folio", "j", "post.ref",
                 "posting ref", "no", "no.", "#", "entry", "entry no", "trans", "transaction"}


@dataclass
class Pair:
    group: str
    debit: int
    credit: int


@dataclass
class Line:
    row: int
    name: str = ""
    name_col: int = 0
    indent: bool = False
    total: bool = False
    has_date: bool = False
    entry: int = 0
    acct_type: str = accounts.UNKNOWN


@dataclass
class AccountTable:
    ws: object
    header_row: int
    acct_cols: list
    pairs: list
    date_col: int = None
    kind: str = "journal"           # journal | tb | worksheet
    title_text: str = ""
    lines: list = field(default_factory=list)
    end_row: int = 0

    @property
    def account_lines(self):
        return [ln for ln in self.lines if not ln.total and ln.name]


@dataclass
class AccountInfo:
    name: str
    acct_type: str
    order: int
    debit_cells: list = field(default_factory=list)    # refs that add to debits
    credit_cells: list = field(default_factory=list)   # refs that add to credits
    debit_ref: str = None                              # balance cells (trial balance)
    credit_ref: str = None

    def dr_minus_cr(self):
        if self.debit_ref or self.credit_ref:
            dr = self.debit_ref or "0"
            cr = self.credit_ref or "0"
            return f"({dr}-{cr})"
        return f"({sum_formula(self.debit_cells)}-{sum_formula(self.credit_cells)})"

    def cr_minus_dr(self):
        if self.debit_ref or self.credit_ref:
            return f"({self.credit_ref or '0'}-{self.debit_ref or '0'})"
        return f"({sum_formula(self.credit_cells)}-{sum_formula(self.debit_cells)})"

    def normal_balance(self):
        """Balance as a positive number on the account's normal side."""
        if accounts.is_debit_normal(self.acct_type) or self.acct_type == accounts.UNKNOWN:
            return self.dr_minus_cr()
        return self.cr_minus_dr()


class AccountBook:
    """Ordered account balances, keyed by account_key()."""

    def __init__(self, source=""):
        self.accounts = {}
        self.source = source

    def get(self, name):
        return self.accounts.get(accounts.account_key(name))

    def add(self, name, acct_type):
        key = accounts.account_key(name)
        if key not in self.accounts:
            self.accounts[key] = AccountInfo(name.strip(), acct_type, len(self.accounts))
        return self.accounts[key]

    def of_type(self, *types):
        return [a for a in self.accounts.values() if a.acct_type in types]

    def __len__(self):
        return len(self.accounts)


# --------------------------------------------------------------------------
# Detection
# --------------------------------------------------------------------------

def _group_label(ctx, ws, header_row, col, lo):
    """Text above a debit column naming its group (Adjustments, Balance Sheet, ...)."""
    for r in (header_row - 1, header_row - 2):
        if r < 1:
            continue
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= r <= rng.max_row and rng.min_col <= col <= rng.max_col:
                v = ws.cell(row=rng.min_row, column=rng.min_col).value
                if is_label(v):
                    return norm(v)
        for c in range(col + 1, lo, -1):
            v = ctx.value(ws, r, c)
            if is_label(v):
                return norm(v)
    return ""


def find_tables(ctx, ws):
    tables = []
    r = 1
    max_row = ws.max_row
    max_col = min(ws.max_column, 40)
    while r <= max_row:
        heads = {}
        for c in range(1, max_col + 1):
            v = ctx.value(ws, r, c)
            if isinstance(v, str):
                heads[c] = norm(v)
        acct = [c for c, t in heads.items() if _is_account_header(t)]
        debits = [c for c, t in heads.items() if _is_debit_header(t)]
        credits = [c for c, t in heads.items() if _is_credit_header(t)]
        if not (debits and credits):
            r += 1
            continue
        if not acct and r > 1:      # "Account Title" one row up (two-row worksheet header)
            acct = [c for c in range(1, min(debits)) if isinstance(ctx.value(ws, r - 1, c), str)
                    and _is_account_header(norm(ctx.value(ws, r - 1, c)))]
        if not acct:                # no header at all: first column holding names below
            for c in range(1, min(debits)):
                if sum(is_label(ctx.value(ws, rr, c)) for rr in range(r + 1, r + 5)) >= 2:
                    acct = [c]
                    break
        if not acct:
            r += 1
            continue
        acct_col = min(acct)
        debits = sorted(c for c in debits if c > acct_col)
        credits = sorted(c for c in credits if c > acct_col)
        pairs = []
        prev = acct_col
        for i, d in enumerate(debits):
            nxt = debits[i + 1] if i + 1 < len(debits) else max_col + 1
            cr = [c for c in credits if d < c < nxt]
            if cr:
                pairs.append(Pair(_group_label(ctx, ws, r, d, prev), d, cr[0]))
                prev = cr[0]
        if not pairs:
            r += 1
            continue
        date_col = next((c for c, t in heads.items() if _is_date_header(t) and c < pairs[0].debit),
                        None)
        skip = {c for c, t in heads.items() if _is_ref_header(t) or _is_date_header(t)}
        acct_cols = [c for c in range(acct_col, pairs[0].debit) if c not in skip]
        title = " ".join(norm(ctx.value(ws, rr, cc)) for rr in range(max(1, r - 4), r)
                         for cc in range(1, max_col + 1) if is_label(ctx.value(ws, rr, cc)))
        table = AccountTable(ws, r, acct_cols, pairs, date_col, title_text=title)
        _parse_lines(ctx, table)
        _decide_kind(ctx, table)
        tables.append(table)
        top = r - 2 if table.kind == "worksheet" else r
        ctx.claim(ws, range(max(1, top), table.end_row + 1))
        r = table.end_row + 1
    return tables


def _parse_lines(ctx, t):
    ws = t.ws
    r = t.header_row + 1
    empty = 0
    entry = 0
    last_was_credit = False
    end = t.header_row
    while r <= ws.max_row:
        name, name_col = "", 0
        for c in t.acct_cols:
            v = ctx.value(ws, r, c)
            if is_label(v):
                name, name_col = v, c
                break
        amounts = [ctx.value(ws, r, c) for p in t.pairs for c in (p.debit, p.credit)]
        has_amount = any(not is_blank(a) for a in amounts)
        date_v = ctx.value(ws, r, t.date_col) if t.date_col else None
        has_date = not is_blank(date_v)
        if not name and not has_amount and not has_date:
            empty += 1
            if empty >= 3:
                break
            if empty == 1:
                entry += 1
            r += 1
            continue
        empty = 0
        n = norm(name)
        if name and not has_amount and TITLE_STOP.match(n) and r > t.header_row + 1:
            break
        if name and _is_account_header(n) and any(
                _is_debit_header(norm(ctx.value(ws, r, p.debit)) if isinstance(
                    ctx.value(ws, r, p.debit), str) else "") for p in t.pairs):
            break
        end = r
        line = Line(row=r, name=str(name).strip(), name_col=name_col, has_date=has_date)
        line.indent = (name_col > t.acct_cols[0]) or (isinstance(name, str) and (
            name.startswith(" ") or name.startswith("\t"))) or bool(
            ws.cell(row=r, column=name_col or 1).alignment.indent) if name else False
        if re.match(r"^(totals?|column totals?)\b", n):
            line.total = True
            t.lines.append(line)
            r += 1
            continue
        if re.match(r"^net (income|loss|profit|earnings)", n):
            r += 1              # worksheet "Net income" row, not an account
            continue
        if name and not has_amount:
            line.acct_type = ctx.classify(name)
            if EXPLANATION.match(n) or line.acct_type == accounts.UNKNOWN:
                r += 1          # explanation text under a journal entry
                continue
        is_debit_line = any(not is_blank(ctx.value(ws, r, p.debit)) for p in t.pairs)
        if has_date or (is_debit_line and last_was_credit and not line.indent):
            entry += 1
        last_was_credit = (not is_debit_line and has_amount) or (line.indent and not has_amount)
        line.entry = entry
        if name:
            line.acct_type = ctx.classify(name)
        t.lines.append(line)
        r += 1
    t.end_row = end


def _decide_kind(ctx, t):
    if len(t.pairs) >= 2:
        t.kind = "worksheet"
        return
    title = t.title_text + " " + norm(t.ws.title)
    names = [accounts.account_key(ln.name) for ln in t.account_lines]
    if "trial balance" in title:
        t.kind = "tb"
    elif "journal" in title:
        t.kind = "journal"
    elif t.date_col and any(ln.has_date for ln in t.lines):
        t.kind = "journal"
    elif len(names) != len(set(names)):
        t.kind = "journal"
    else:
        t.kind = "tb"


# --------------------------------------------------------------------------
# Filling missing amounts
# --------------------------------------------------------------------------

def _num(ctx, ws, r, c):
    return ctx.number(ws, r, c)


def fill_missing(ctx, t):
    if t.kind == "journal":
        _fill_journal(ctx, t)
    elif t.kind == "tb":
        _fill_tb(ctx, t)
    else:
        from .worksheet10 import fill_worksheet
        fill_worksheet(ctx, t)
    _fill_totals(ctx, t)


def _fill_journal(ctx, t):
    ws = t.ws
    p = t.pairs[0]
    by_entry = {}
    for ln in t.account_lines:
        by_entry.setdefault(ln.entry, []).append(ln)
    for entry, lines in by_entry.items():
        missing = [ln for ln in lines if ctx.is_empty_cell(ws, ln.row, p.debit)
                   and ctx.is_empty_cell(ws, ln.row, p.credit)]
        known = [ln for ln in lines if ln not in missing]
        if len(missing) != 1 or not known:
            if len(missing) > 1:
                ctx.warn(f"{ws.title}: journal entry starting row {lines[0].row} has "
                         f"{len(missing)} lines with no amount - Solver can only work out one.")
            continue
        ln = missing[0]
        has_dr = any(not ctx.is_empty_cell(ws, k.row, p.debit) for k in known)
        has_cr = any(not ctx.is_empty_cell(ws, k.row, p.credit) for k in known)
        if ln.indent or not has_cr:
            side = "credit"
        elif not has_dr:
            side = "debit"
        else:
            side = "debit"
        dr_refs = [cell_ref(ws, k.row, p.debit, ws) for k in known]
        cr_refs = [cell_ref(ws, k.row, p.credit, ws) for k in known]
        if side == "credit":
            formula = f"={sum_formula(dr_refs)}-{sum_formula(cr_refs)}"
            col = p.credit
        else:
            formula = f"={sum_formula(cr_refs)}-{sum_formula(dr_refs)}"
            col = p.debit
        ctx.fill(ws, ln.row, col, formula,
                 f"Missing {side} for {ln.name}: debits must equal credits in the entry")


def _fill_tb(ctx, t):
    ws = t.ws
    p = t.pairs[0]
    lines = t.account_lines
    missing = [ln for ln in lines if ctx.is_empty_cell(ws, ln.row, p.debit)
               and ctx.is_empty_cell(ws, ln.row, p.credit)]
    if not missing:
        return
    if len(missing) > 1:
        ctx.warn(f"{ws.title}: {len(missing)} trial balance accounts have no balance; Solver can "
                 f"only work out a single missing balance.")
        return
    ln = missing[0]
    others = [k for k in lines if k is not ln]
    total = next((k for k in t.lines if k.total), None)
    debit_side = accounts.is_debit_normal(ln.acct_type)
    col = p.debit if debit_side else p.credit
    other_col = p.credit if debit_side else p.debit
    same = [cell_ref(ws, k.row, col, ws) for k in others]
    if total and not ctx.is_empty_cell(ws, total.row, col) and \
            ctx.value(ws, total.row, col) is not FORMULA and \
            not str(ws.cell(row=total.row, column=col).value).startswith("="):
        formula = f"={cell_ref(ws, total.row, col, ws)}-{sum_formula(same)}"
        why = "column total minus the other balances"
    else:
        opp = [cell_ref(ws, k.row, other_col, ws) for k in others]
        formula = f"={sum_formula(opp)}-{sum_formula(same)}"
        why = "the amount that makes total debits equal total credits"
    ctx.fill(ws, ln.row, col, formula, f"Missing balance for {ln.name}: {why}")


def _fill_totals(ctx, t):
    ws = t.ws
    first = t.header_row + 1
    for ln in t.lines:
        if not ln.total:
            continue
        rows_above = [k.row for k in t.account_lines if k.row < ln.row]
        if not rows_above:
            continue
        for p in t.pairs:
            for col in (p.debit, p.credit):
                if ctx.is_empty_cell(ws, ln.row, col):
                    letter = get_column_letter(col)
                    ctx.fill(ws, ln.row, col, f"=SUM({letter}{first}:{letter}{ln.row - 1})",
                             "Column total")
        break       # later total rows (worksheet "Net income" totals) are handled elsewhere


# --------------------------------------------------------------------------
# Account balances
# --------------------------------------------------------------------------

def book_from_journal(ctx, tables):
    """Collect postings from one or more journal tables."""
    book = AccountBook("journal")
    for t in tables:
        p = t.pairs[0]
        for ln in t.account_lines:
            info = book.add(ln.name, ln.acct_type)
            info.debit_cells.append(cell_ref(t.ws, ln.row, p.debit))
            info.credit_cells.append(cell_ref(t.ws, ln.row, p.credit))
    return book


def book_from_balances(ctx, t, pair):
    """Balances read from a trial balance (or the adjusted columns of a worksheet)."""
    book = AccountBook(f"{t.ws.title}")
    for ln in t.account_lines:
        info = book.add(ln.name, ln.acct_type)
        info.debit_ref = cell_ref(t.ws, ln.row, pair.debit)
        info.credit_ref = cell_ref(t.ws, ln.row, pair.credit)
    return book
