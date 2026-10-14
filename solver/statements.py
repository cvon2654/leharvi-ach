"""Build the trial balance, general ledger and financial statements as new sheets."""

import datetime as _dt
import re

from openpyxl.utils import get_column_letter

from . import accounts as A
from .workbook import (ACCOUNTING_FMT, BOLD, NOTE_FONT, TITLE, cell_ref, is_label, norm, set_widths,
                       style_header, sum_formula, underline, write_title)

TYPE_ORDER = [A.ASSET, A.CONTRA_ASSET, A.LIABILITY, A.EQUITY, A.CONTRA_EQUITY, A.DRAWING,
              A.REVENUE, A.CONTRA_REVENUE, A.COGS, A.CONTRA_COGS, A.EXPENSE, A.CLOSING,
              A.UNKNOWN]

IS_TITLES = re.compile(r"income statement|statement of (comprehensive )?income|profit and loss|"
                       r"statement of operations|statement of earnings")
EQ_TITLES = re.compile(r"statement of (retained earnings|owner'?s'? equity|stockholders'? equity|"
                       r"shareholders'? equity|changes in (owner'?s'? )?equity|partners'? equity|"
                       r"members'? equity)|retained earnings statement")
BS_TITLES = re.compile(r"balance sheet|statement of financial position")
SKIP_TITLE = re.compile(r"journal|trial balance|ledger|worksheet|work sheet|problem|exercise|"
                        r"instructions|required|chapter|^date|^account|^for the|^as of|"
                        r"^december|^january|statement|balance sheet|^\d")


def ordered(accts):
    return sorted(accts, key=lambda a: (TYPE_ORDER.index(a.acct_type), a.order))


def guess_company(ctx):
    for ws in ctx.wb.worksheets:
        if ws.title in ctx.created:
            continue
        for r in range(1, 4):
            if ctx.is_claimed(ws, r):
                break           # reached a table header: no company name above it
            for c in range(1, 6):
                v = ctx.value(ws, r, c)
                if is_label(v) and not SKIP_TITLE.search(norm(v)) and len(v) < 60:
                    return v.strip()
    return None


def guess_period_end(ctx, tables):
    dates = []
    for t in tables:
        if t.date_col:
            for ln in t.lines:
                v = ctx.value(t.ws, ln.row, t.date_col)
                if isinstance(v, (_dt.date, _dt.datetime)):
                    dates.append(v)
    if not dates:
        return None
    d = max(dates)
    return f"{d:%B} {d.day}, {d.year}"


def existing_statements(ctx):
    """Which statements the workbook already has (as templates to fill)."""
    found = set()
    for ws in ctx.wb.worksheets:
        if ws.title in ctx.created:
            continue
        texts = [norm(ws.title)]
        labels = set()
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400),
                                max_col=min(ws.max_column, 12)):
            if row and ctx.is_claimed(ws, row[0].row):
                continue        # e.g. "Income Statement" columns of a worksheet
            for cell in row:
                if is_label(cell.value):
                    t = norm(cell.value)
                    texts.append(t)
                    labels.add(t)
        for t in texts:
            if IS_TITLES.search(t):
                found.add("is")
            if EQ_TITLES.search(t):
                found.add("equity")
            if BS_TITLES.search(t):
                found.add("bs")
        if any(t.startswith("total assets") for t in labels) and \
                any(t.startswith("total liabilities") for t in labels):
            found.add("bs")
    return found


class Writer:
    """Writes a two-amount-column statement: items in B, totals in C."""

    def __init__(self, ctx, ws):
        self.ctx = ctx
        self.ws = ws
        self.row = 1

    def titles(self, lines):
        self.row = write_title(self.ws, lines)

    def header(self, text, indent=0):
        c = self.ws.cell(row=self.row, column=1, value=("    " * indent) + text)
        c.font = BOLD
        self.row += 1

    def line(self, label, formula, col, bold=False, indent=1, rule=None):
        ws = self.ws
        c = ws.cell(row=self.row, column=1, value=("    " * indent) + label)
        if bold:
            c.font = BOLD
        v = ws.cell(row=self.row, column=col, value=formula)
        v.number_format = ACCOUNTING_FMT
        if bold:
            v.font = BOLD
        if rule:
            underline(ws, self.row, col, double=(rule == "double"))
        ref = f"{get_column_letter(col)}{self.row}"
        self.row += 1
        return ref

    def item(self, label, formula, indent=1):
        return self.line(label, formula, 2, indent=indent)

    def total(self, label, formula, rule="single", indent=0, bold=True):
        return self.line(label, formula, 3, bold=bold, indent=indent, rule=rule)

    def blank(self):
        self.row += 1

    def full(self, ref):
        return f"{_q(self.ws.title)}!{ref}"


def _q(title):
    from .workbook import quote_sheet
    return quote_sheet(title)


def _sum(refs):
    return "=" + sum_formula(refs) if refs else "=0"


# --------------------------------------------------------------------------

def create_trial_balance(ctx, book, journals, company, period):
    ws = ctx.new_sheet("Trial Balance")
    row = write_title(ws, [company, "Trial Balance", period and f"{period}"])
    ws.cell(row=row - 1, column=1,
            value=f"Built by Solver from: {', '.join(sorted({t.ws.title for t in journals}))}"
            ).font = NOTE_FONT
    hdr = row
    for col, text in enumerate(["Account", "Debit", "Credit", "", "Posted debits",
                                "Posted credits"], start=1):
        ws.cell(row=hdr, column=col, value=text or None)
    style_header(ws, hdr, [1, 2, 3, 5, 6])
    r = hdr + 1
    first = r
    for a in ordered(book.accounts.values()):
        ws.cell(row=r, column=1, value=a.name)
        ws.cell(row=r, column=5, value="=" + sum_formula(a.debit_cells))
        ws.cell(row=r, column=6, value="=" + sum_formula(a.credit_cells))
        ws.cell(row=r, column=2, value=f"=MAX(0,E{r}-F{r})")
        ws.cell(row=r, column=3, value=f"=MAX(0,F{r}-E{r})")
        for c in (2, 3, 5, 6):
            ws.cell(row=r, column=c).number_format = ACCOUNTING_FMT
        a.debit_ref = cell_ref(ws, r, 2)
        a.credit_ref = cell_ref(ws, r, 3)
        r += 1
    ws.cell(row=r, column=1, value="Totals").font = BOLD
    for c in (2, 3, 5, 6):
        L = get_column_letter(c)
        cell = ws.cell(row=r, column=c, value=f"=SUM({L}{first}:{L}{r - 1})")
        cell.number_format = ACCOUNTING_FMT
        cell.font = BOLD
        underline(ws, r, c, double=True)
    ws.cell(row=r + 2, column=1, value="Difference (should be 0)")
    ws.cell(row=r + 2, column=2, value=f"=B{r}-C{r}").number_format = ACCOUNTING_FMT
    set_widths(ws, [38, 16, 16, 3, 16, 16])
    ctx.info(f"Created sheet '{ws.title}' with {len(book)} accounts (formulas link back to "
             f"the journal, so they update if you change an entry).")
    return ws


def create_ledger(ctx, book, journals):
    ws = ctx.new_sheet("General Ledger")
    row = write_title(ws, ["General Ledger", "Each account's postings with a running balance"])
    postings = {k: [] for k in book.accounts}
    for t in journals:
        p = t.pairs[0]
        date_row = None
        last_entry = None
        for ln in t.lines:
            if ln.entry != last_entry:
                date_row = None
                last_entry = ln.entry
            if ln.has_date:
                date_row = ln.row
            if ln.total or not ln.name:
                continue
            key = A.account_key(ln.name)
            postings[key].append((t, ln, date_row, p))
    for a in ordered(book.accounts.values()):
        key = A.account_key(a.name)
        normal = "debit" if A.is_debit_normal(a.acct_type) else "credit"
        c = ws.cell(row=row, column=1, value=f"{a.name}  ({a.acct_type}, normal {normal} balance)")
        c.font = BOLD
        row += 1
        for col, text in enumerate(["Date", "Journal row", "Debit", "Credit", "Balance"], 1):
            ws.cell(row=row, column=col, value=text)
        style_header(ws, row, range(1, 6))
        row += 1
        prev = None
        for t, ln, date_row, p in postings[key]:
            if date_row and t.date_col:
                d = ws.cell(row=row, column=1, value="=" + cell_ref(t.ws, date_row, t.date_col))
                d.number_format = t.ws.cell(row=date_row, column=t.date_col).number_format
            ws.cell(row=row, column=2, value=f"{t.ws.title} row {ln.row}")
            ws.cell(row=row, column=3, value="=" + cell_ref(t.ws, ln.row, p.debit))
            ws.cell(row=row, column=4, value="=" + cell_ref(t.ws, ln.row, p.credit))
            change = f"C{row}-D{row}" if normal == "debit" else f"D{row}-C{row}"
            ws.cell(row=row, column=5, value=f"={prev}+{change}" if prev else f"={change}")
            for col in (3, 4, 5):
                ws.cell(row=row, column=col).number_format = ACCOUNTING_FMT
            prev = f"E{row}"
            row += 1
        row += 1
    set_widths(ws, [14, 22, 15, 15, 15])
    ctx.info(f"Created sheet '{ws.title}' (postings for every account).")
    return ws


def _is_corporation(book):
    for a in book.of_type(A.EQUITY, A.CONTRA_EQUITY, A.DRAWING):
        n = a.name.lower()
        if any(w in n for w in ("stock", "retained earnings", "dividend", "share", "paid-in",
                                "paid in")):
            return True
    return False


def _closing_equity_account(book):
    eq = book.of_type(A.EQUITY)
    for a in eq:
        if "retained" in a.name.lower():
            return a
    for a in eq:
        if "capital" in a.name.lower() and "paid" not in a.name.lower() and \
                "stock" not in a.name.lower():
            return a
    return None


def create_income_statement(ctx, book, company, period):
    rev = ordered(book.of_type(A.REVENUE))
    contra = ordered(book.of_type(A.CONTRA_REVENUE))
    cogs = ordered(book.of_type(A.COGS, A.CONTRA_COGS))
    exp = ordered(book.of_type(A.EXPENSE))
    if not (rev or exp or cogs):
        return None
    ws = ctx.new_sheet("Income Statement")
    w = Writer(ctx, ws)
    w.titles([company, "Income Statement",
              f"For the Period Ended {period}" if period else None])
    w.header("Revenues")
    refs = [w.item(a.name, "=" + a.cr_minus_dr()) for a in rev]
    refs += [w.item("Less: " + a.name, "=" + a.cr_minus_dr()) for a in contra]
    rev_label = "Net sales" if (contra or cogs) else "Total revenues"
    top = w.total(rev_label, _sum(refs))
    ctx.register("total revenues", w.full(top), ws)
    ctx.register("net sales", w.full(top), ws)
    if cogs:
        if len(cogs) == 1:
            c = w.total("Cost of goods sold", "=" + cogs[0].dr_minus_cr(), rule=None, bold=False)
        else:
            w.header("Cost of goods sold")
            crefs = [w.item(a.name, "=" + a.dr_minus_cr()) for a in cogs]
            c = w.total("Total cost of goods sold", _sum(crefs))
        if any("purchases" in a.name.lower() for a in cogs):
            ctx.warn("Purchases account found (periodic inventory). Cost of goods sold also needs "
                     "beginning and ending inventory - adjust the Income Statement if your "
                     "problem gives an ending inventory count.")
        ctx.register("cost of goods sold", w.full(c), ws)
        top = w.total("Gross profit", f"={top}-{c}")
        ctx.register("gross profit", w.full(top), ws)
    w.blank()
    erefs = []
    if exp:
        w.header("Operating expenses" if cogs else "Expenses")
        erefs = [w.item(a.name, "=" + a.dr_minus_cr()) for a in exp]
        etot = w.total("Total operating expenses" if cogs else "Total expenses", _sum(erefs))
        ctx.register("total expenses", w.full(etot), ws)
        ni = w.total("Net income (loss)", f"={top}-{etot}", rule="double")
    else:
        ni = w.total("Net income (loss)", f"={top}", rule="double")
    ctx.register("net income", w.full(ni), ws, overwrite=True)
    set_widths(ws, [42, 16, 16])
    ctx.info(f"Created sheet '{ws.title}'.")
    return ws


def create_equity_statement(ctx, book, company, period):
    corp = _is_corporation(book)
    closing = _closing_equity_account(book)
    draws = ordered(book.of_type(A.DRAWING))
    ni = ctx.lookup("net income")
    if not ni and not draws:
        return None
    title = "Statement of Retained Earnings" if corp else "Statement of Owner's Equity"
    ws = ctx.new_sheet("Retained Earnings" if corp else "Owner's Equity")
    w = Writer(ctx, ws)
    w.titles([company, title, f"For the Period Ended {period}" if period else None])
    base = "Retained earnings" if corp else (closing.name if closing else "Owner's capital")
    refs = [w.total(f"{base}, beginning", "=" + closing.cr_minus_dr() if closing else "=0",
                    rule=None, bold=False)]
    if not corp:
        for a in book.of_type(A.EQUITY):
            if a is not closing:
                refs.append(w.total(f"Add: {a.name}", "=" + a.cr_minus_dr(), rule=None,
                                    bold=False))
    if ni:
        refs.append(w.total("Add: Net income (loss)", "=" + ni, rule=None, bold=False))
    for a in draws:
        refs.append(w.total(f"Less: {a.name}", "=-" + a.dr_minus_cr(), rule=None, bold=False))
    end = w.total(f"{base}, ending", _sum(refs), rule="double")
    key = "ending retained earnings" if corp else "ending capital"
    ctx.register(key, w.full(end), ws, overwrite=True)
    ctx.register("ending equity", w.full(end), ws, overwrite=True)
    set_widths(ws, [42, 4, 16])
    ctx.info(f"Created sheet '{ws.title}'.")
    return ws


def create_balance_sheet(ctx, book, company, period):
    assets = ordered(book.of_type(A.ASSET, A.CONTRA_ASSET))
    liabs = ordered(book.of_type(A.LIABILITY))
    if not assets:
        return None
    corp = _is_corporation(book)
    closing = _closing_equity_account(book)
    ws = ctx.new_sheet("Balance Sheet")
    w = Writer(ctx, ws)
    w.titles([company, "Balance Sheet", period])

    def section(title, accts, total_label, value):
        if not accts:
            return None
        w.header(title, indent=0 if title in ("Assets",) else 1)
        refs = [w.item(("Less: " if a.acct_type in (A.CONTRA_ASSET, A.CONTRA_EQUITY) else "")
                       + a.name, "=" + value(a), indent=2) for a in accts]
        return w.total(total_label, _sum(refs), indent=1)

    w.header("Assets")
    cur = [a for a in assets if A.is_current(a.name, a.acct_type)]
    lt = [a for a in assets if a not in cur]
    subs = [s for s in (
        section("Current assets", cur, "Total current assets", lambda a: a.dr_minus_cr()),
        section("Property, plant and equipment / long-term assets", lt,
                "Total long-term assets", lambda a: a.dr_minus_cr())) if s]
    if cur:
        ctx.register("total current assets", w.full(subs[0]), ws)
    ta = w.total("Total assets", _sum(subs), rule="double")
    ctx.register("total assets", w.full(ta), ws)
    w.blank()

    w.header("Liabilities")
    cl = [a for a in liabs if A.is_current(a.name, a.acct_type)]
    ll = [a for a in liabs if a not in cl]
    lsubs = []
    s = section("Current liabilities", cl, "Total current liabilities", lambda a: a.cr_minus_dr())
    if s:
        lsubs.append(s)
        ctx.register("total current liabilities", w.full(s), ws)
    s = section("Long-term liabilities", ll, "Total long-term liabilities",
                lambda a: a.cr_minus_dr())
    if s:
        lsubs.append(s)
    tl = w.total("Total liabilities", _sum(lsubs))
    ctx.register("total liabilities", w.full(tl), ws)
    w.blank()

    eq_title = "Stockholders' equity" if corp else "Owner's equity"
    w.header(eq_title)
    erefs = []
    for a in ordered(book.of_type(A.EQUITY, A.CONTRA_EQUITY)):
        if a is closing:
            continue
        label = ("Less: " if a.acct_type == A.CONTRA_EQUITY else "") + a.name
        erefs.append(w.item(label, "=" + a.cr_minus_dr(), indent=2))
    ending = ctx.lookup("ending retained earnings", "ending capital", "ending equity")
    label = "Retained earnings" if corp else (closing.name if closing else "Owner's capital")
    if ending:
        erefs.append(w.item(label, "=" + ending, indent=2))
    else:
        # No equity statement: closing account + net income - drawings, in one formula.
        parts = [closing.cr_minus_dr()] if closing else []
        ni = ctx.lookup("net income")
        if ni:
            parts.append(ni)
        parts += ["-" + d.dr_minus_cr() for d in book.of_type(A.DRAWING)]
        if parts:
            erefs.append(w.item(label, "=" + "+".join(parts).replace("+-", "-"), indent=2))
    te = w.total(f"Total {eq_title.lower()}", _sum(erefs), indent=1)
    ctx.register("total equity", w.full(te), ws)
    tle = w.total(f"Total liabilities and {eq_title.lower()}", f"={tl}+{te}", rule="double")
    ctx.register("total liabilities and equity", w.full(tle), ws)
    w.blank()
    chk = w.line("Check: total assets minus total liabilities and equity (should be 0)",
                 f"={ta}-{tle}", 3, indent=0)
    ws[chk].font = NOTE_FONT
    unknown = book.of_type(A.UNKNOWN)
    if unknown:
        ctx.warn("Could not classify these accounts, so they are left off the statements: "
                 + ", ".join(a.name for a in unknown) +
                 ". Add a 'Chart of Accounts' sheet (Account | Type) to tell Solver.")
    set_widths(ws, [52, 16, 16])
    ctx.info(f"Created sheet '{ws.title}'.")
    return ws
