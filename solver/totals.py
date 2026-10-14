"""Fill blanks in statements the workbook already contains.

A template is a column of labels with amounts to the right.  For every
blank amount Solver either
  * links it to a known figure (an account balance, or a total from another
    sheet such as Net income), or
  * computes it from the lines above it when the label is a total
    ("Total current assets", "Gross profit", "Net income", "Ending
    retained earnings", ...).
"""

import re
from collections import Counter
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter

from . import accounts as A
from . import ratios
from .workbook import (FORMULA, cell_ref, is_blank, is_label, is_number, is_placeholder, norm)

MAX_ROWS = 600
MAX_LABEL_COL = 8

SECTION_HEADERS = {
    "revenues", "revenue", "revenues and gains", "other revenues and gains", "expenses",
    "operating expenses", "other expenses", "other expenses and losses", "expenses and losses",
    "selling expenses", "administrative expenses", "general and administrative expenses",
    "assets", "current assets", "long term assets", "noncurrent assets", "non current assets",
    "property plant and equipment", "property, plant and equipment", "plant assets",
    "long term investments", "investments", "intangible assets", "other assets",
    "liabilities", "current liabilities", "long term liabilities", "noncurrent liabilities",
    "non current liabilities", "stockholders' equity", "stockholders equity",
    "shareholders' equity", "shareholders equity", "owner's equity", "owners equity", "equity",
    "liabilities and stockholders' equity", "liabilities and stockholders equity",
    "liabilities and owner's equity", "liabilities and owners equity",
    "liabilities and equity", "liabilities and shareholders' equity", "other income",
    "other revenues", "other income and expense", "other income (expense)",
    "other revenues and expenses", "paid-in capital", "paid in capital", "capital stock",
    "nonoperating items", "non-operating items", "sales", "cost of goods sold",
}

TITLE_RE = re.compile(r"^(statement of|income statement|balance sheet|for the (year|month|period|"
                      r"quarter)|(year|month|period) ended|as of\b|schedule of)")
TOTAL_RE = re.compile(
    r"^(total|subtotal|sub-total|net |gross (profit|margin)|income (before|from)|"
    r"operating (income|profit|loss)|(income|profit|loss) from operations|"
    r"cost of goods (sold|available)|goods available|ending)")
ENDING_EQUITY_RE = re.compile(r"(retained earnings|capital|equity).*(\bend(ing)?\b|end of)|"
                              r"^ending (retained earnings|capital|equity)")
BEGIN_EQUITY_RE = re.compile(r"(retained earnings|capital|equity).*(begin|start)|"
                             r"^beginning (retained earnings|capital|equity)")
NEGATIVE_WORDS = re.compile(r"\b(less|deduct|minus)\b|accumulated depreciation|allowance|"
                            r"returns|discounts|drawing|withdrawal|dividend|treasury|net loss|"
                            r"\bloss\b|losses")
EXPENSE_WORDS = re.compile(r"expense|cost of goods|cost of sales|cogs|\btax|loss|"
                           r"returns|discounts|allowances")


def canonical(label):
    """A concept key for a label: "Net Income (Loss)" -> "net income"."""
    t = norm(label)
    t = re.sub(r"^(add|less|plus|deduct|minus)\b\s*:?\s*", "", t)
    t = re.sub(r"\(.*?\)", " ", t)
    t = t.replace("'", "")
    t = re.sub(r"\s+", " ", t).strip(" ,.")
    retained = "retained earnings" in t
    if re.search(r"\bbefore\b.*\btax", t):
        return "income before taxes"
    if re.search(r"\bnet (income|loss|earnings|profit)\b", t) and "ratio" not in t and \
            "per share" not in t and "margin" not in t:
        return "net income"
    if re.search(r"gross (profit|margin)", t) and not re.search(r"rate|ratio|percent|%", t):
        return "gross profit"
    if re.fullmatch(r"(total )?net (sales|revenues?)( revenues?)?", t):
        return "net sales"
    if re.search(r"(income|profit|loss) from operations|operating (income|profit|loss)", t):
        return "operating income"
    if re.search(r"cost of goods available|goods available for sale", t):
        return "cost of goods available for sale"
    if re.fullmatch(r"(total )?cost of (goods sold|sales|merchandise sold)", t):
        return "cost of goods sold"
    if ENDING_EQUITY_RE.search(t):
        return "ending retained earnings" if retained else "ending capital"
    if BEGIN_EQUITY_RE.search(t):
        return "beginning retained earnings" if retained else "beginning capital"
    m = re.match(r"^total (.*)$", t)
    if m:
        rest = m.group(1)
        rest = re.sub(r"\b(stockholders|shareholders|owners|partners|members) equity", "equity",
                      rest)
        rest = re.sub(r"\bowner equity\b", "equity", rest)
        rest = re.sub(r"\brevenue\b", "revenues", rest)
        rest = re.sub(r"^(liabilities and|liabilities &) .*equity$", "liabilities and equity",
                      rest)
        rest = re.sub(r"^(liabilities and|liabilities &) capital$", "liabilities and equity",
                      rest)
        if rest in ("sales", "revenues and gains"):
            rest = "revenues"
        if rest in ("expenses and losses",):
            rest = "expenses"
        return "total " + rest
    return t


def is_total_label(label):
    t = norm(label)
    t = re.sub(r"^(add|less|plus|deduct|minus)\b\s*:?\s*", "", t)
    if TOTAL_RE.match(t) or ENDING_EQUITY_RE.search(t):
        return True
    if t in ("net sales", "net revenues", "net revenue"):
        return True
    return False


def _source(expr, here):
    """'Trial Balance' for an expression pointing at that sheet, for the explanation."""
    m = re.search(r"'((?:[^']|'')+)'!|([A-Za-z_][\w.]*)!", expr)
    sheet = (m.group(1) or m.group(2)).replace("''", "'") if m else here
    return f"from {sheet}" if sheet != here else "figure on this sheet"


def _fits_under(part, whole):
    """May total `part` be a piece of total `whole`?  ("total current assets" fits under
    "total assets"; "total liabilities" does not fit under "total equity".)"""
    if not (part.startswith("total ") and whole.startswith("total ")):
        return True
    p, w = part[6:], whole[6:]
    heads = [x.strip() for x in re.split(r"\band\b|&", w)]
    return any(h and (h in p or p in h) for h in heads)


def _bs_side(label):
    t = norm(label)
    if "liabilit" in t or "equity" in t or "capital" in t or "retained" in t:
        return "LE"
    if "asset" in t:
        return "A"
    return None


@dataclass
class Row:
    r: int
    label: str
    label_col: int
    nums: dict = field(default_factory=dict)        # col -> True for cells holding values
    holes: list = field(default_factory=list)       # placeholder columns ("?" / "____")
    refill: list = field(default_factory=list)      # columns Solver filled on an earlier pass
    kind: str = "item"                              # item | total | header | wall
    value_col: int = None                           # where this row's amount is/will be
    cover: int = None                               # first row covered by this total


class SheetFiller:
    def __init__(self, ctx, ws, final):
        self.ctx = ctx
        self.ws = ws
        self.final = final
        self.rows = {}
        self.amount_cols = []
        self.comparative = False
        self.title_kind = {}    # row -> "is" | "bs" | "equity"

    # ------------------------------------------------------------------
    def _cell_has_value(self, r, c):
        if (self.ws.title, f"{get_column_letter(c)}{r}") in self.ctx.template_fills:
            return False        # our own answer from an earlier pass: recompute it
        v = self.ctx.value(self.ws, r, c)
        if v is FORMULA or is_number(v):
            return True
        if isinstance(v, str) and self.ctx.number(self.ws, r, c) is not None:
            return True
        return False

    def _hole(self, r, c):
        return is_placeholder(self.ws.cell(row=r, column=c).value)

    def _refill(self, r, c):
        return (self.ws.title, f"{get_column_letter(c)}{r}") in self.ctx.template_fills

    def scan(self):
        ws, ctx = self.ws, self.ctx
        max_col = min(ws.max_column, 20)
        kind = None
        for r in range(1, min(ws.max_row, MAX_ROWS) + 1):
            if ctx.is_claimed(ws, r):
                self.rows[r] = Row(r, "", 0, kind="wall")
                continue
            label, lc = None, 0
            for c in range(1, min(max_col, MAX_LABEL_COL) + 1):
                v = ctx.value(ws, r, c)
                if is_label(v):
                    label, lc = v, c
                    break
            if label is None:
                continue
            t = norm(label)
            if re.search(r"income statement|statement of (comprehensive )?income|"
                         r"statement of operations|profit and loss", t):
                kind = "is"
            elif re.search(r"retained earnings statement|statement of (retained|owner|"
                           r"stockholder|shareholder|changes|partner)", t):
                kind = "equity"
            elif re.search(r"balance sheet|financial position", t):
                kind = "bs"
            self.title_kind[r] = kind
            row = Row(r, label.strip(), lc)
            for c in range(lc + 1, max_col + 1):
                if self._cell_has_value(r, c):
                    row.nums[c] = True
                elif self._refill(r, c):
                    row.refill.append(c)
                elif self._hole(r, c):
                    row.holes.append(c)
            self.rows[r] = row
        cols = Counter()
        for row in self.rows.values():
            for c in list(row.nums) + row.holes + row.refill:
                cols[c] += 1
        self.amount_cols = sorted(cols)
        self._detect_comparative()

    def _detect_comparative(self):
        """Two or more year columns (2025 | 2024): every column is its own statement."""
        if len(self.amount_cols) < 2:
            return
        first = min((r for r, row in self.rows.items() if row.nums), default=None)
        if first is None:
            return
        hits = 0
        for c in self.amount_cols:
            for r in range(max(1, first - 4), first):
                v = self.ctx.value(self.ws, r, c)
                if (is_number(v) and 1950 <= v <= 2100) or (
                        isinstance(v, str) and re.search(r"\b(19|20)\d\d\b|current|prior|"
                                                         r"this year|last year", v.lower())):
                    hits += 1
                    break
        self.comparative = hits >= 2

    # ------------------------------------------------------------------
    def _labels_above(self, r):
        out = []
        for rr in sorted((x for x in self.rows if x < r), reverse=True):
            other = self.rows[rr]
            if other.kind == "wall":
                break
            t = norm(other.label)
            if t in SECTION_HEADERS and out:
                out.append(t)
                break
            if self.title_kind.get(rr) != self.title_kind.get(r):
                break
            out.append(t)
        return " | ".join(out)

    def _computable(self, row):
        """Is this total-looking label really computed from the lines above it?"""
        key = canonical(row.label)
        above = self._labels_above(row.r)
        if key in ("net income", "income before taxes", "operating income"):
            return bool(re.search(r"expense|revenue|sales|cost|gross|income before|fees|"
                                  r"operating", above))
        if key == "gross profit":
            return "cost" in above and bool(re.search(r"sales|revenue", above))
        if key == "cost of goods sold":
            return bool(re.search(r"inventory|purchases|goods available", above))
        if key == "net sales":
            return bool(re.search(r"return|discount|allowance", above))
        if key == "cost of goods available for sale":
            return bool(re.search(r"inventory|purchases", above))
        return bool(above)

    def _in_statement(self, row):
        """Only fill item blanks inside something that looks like a statement/table."""
        if len(row.label) > 70 or len(row.label.split()) > 9:
            return False
        if row.nums or row.holes or row.refill or self.statement_kind(row.r):
            return True
        near = [self.rows[r] for r in self.rows if abs(r - row.r) <= 3 and r != row.r]
        return any(o.nums or o.holes for o in near)

    def _equity_positions(self):
        """'Retained earnings, Jan 1' ... 'Retained earnings, Dec 31': first is beginning,
        last is ending, when net income / dividends lines sit between them."""
        self.key_override = {}
        rows = [row for r, row in sorted(self.rows.items()) if row.kind != "wall" and
                re.search(r"retained earnings|capital|equity", norm(row.label)) and
                not re.search(r"statement|total|paid|stock|liabilit", norm(row.label))]
        for a, b in zip(rows, rows[1:]):
            if self.title_kind.get(a.r) != self.title_kind.get(b.r):
                continue
            between = " ".join(norm(self.rows[r].label) for r in self.rows if a.r < r < b.r)
            if re.search(r"net income|net loss|dividend|drawing|withdraw", between):
                word = "retained earnings" if "retained" in norm(a.label) else "capital"
                self.key_override.setdefault(a.r, f"beginning {word}")
                self.key_override[b.r] = f"ending {word}"

    def key(self, row):
        return self.key_override.get(row.r) or canonical(row.label)

    def classify_rows(self):
        self._equity_positions()
        for row in self.rows.values():
            if row.kind == "wall":
                continue
            t = norm(row.label)
            if not self._in_statement(row):
                row.kind = "header"
            elif TITLE_RE.search(t) and not (row.nums or row.holes):
                row.kind = "header"
            elif self.key(row).startswith("ending ") and self.key_override.get(row.r):
                row.kind = "total"
            elif ratios.is_ratio(row.label):
                row.kind = "item"
            elif is_total_label(row.label) and (self._computable(row) or (
                    not row.nums and (row.holes or not self.link(row)))) and \
                    not self.key_override.get(row.r, "").startswith("beginning"):
                row.kind = "total"
            elif row.nums or row.holes or row.refill:
                row.kind = "item"
            elif t in SECTION_HEADERS or t.rstrip(":") in SECTION_HEADERS:
                row.kind = "header"
            elif re.match(r"^(add|less|plus|deduct)\b", t) or self.ctx.classify(row.label) != \
                    A.UNKNOWN or self.link(row) or ratios.is_ratio(row.label):
                row.kind = "item"
            else:
                row.kind = "header"

    def statement_kind(self, r):
        k = self.title_kind.get(r)
        if k:
            return k
        labels = " ".join(norm(row.label) for row in self.rows.values())
        if "total assets" in labels:
            return "bs"
        if "net income" in labels and ("revenue" in labels or "expense" in labels):
            return "is"
        return None

    # ------------------------------------------------------------------
    def link(self, row):
        """Expression for an item label that names a known figure, else None."""
        ctx = self.ctx
        key = self.key(row) if hasattr(self, "key_override") else canonical(row.label)
        kind = self.statement_kind(row.r)
        book = ctx.book
        if ratios.is_ratio(row.label):
            return None
        if key in ("ending retained earnings", "ending capital"):
            got = ctx.lookup(key, "ending equity")
            if got and ctx.registry_src.get(key) != self.ws.title:
                return got
        if key in ("beginning retained earnings", "beginning capital") and book is not None:
            eq = [a for a in book.of_type(A.EQUITY)
                  if ("retained" in a.name.lower()) == (key == "beginning retained earnings")]
            if eq:
                return eq[0].normal_balance()
        plain = re.sub(r"^(add|less|plus|deduct|minus)\b\s*:?\s*", "", norm(row.label))
        acct = None
        if book is not None:
            acct = book.get(plain) or self._fuzzy_account(plain)
        if acct is not None:
            if kind == "bs" and acct.acct_type == A.EQUITY and (
                    "retained" in acct.name.lower() or "capital" in acct.name.lower()) and \
                    "paid" not in acct.name.lower():
                ending = ctx.lookup("ending retained earnings", "ending capital", "ending equity")
                if ending:
                    return ending
                if not self.final and book.of_type(A.REVENUE, A.EXPENSE, A.DRAWING):
                    return None         # wait for the equity statement / net income
            return acct.normal_balance()
        if key in ("beginning retained earnings", "beginning capital") and book is not None and \
                acct is None:
            eq = [a for a in book.of_type(A.EQUITY)
                  if ("retained" in a.name.lower()) == (key == "beginning retained earnings")]
            if eq:
                return eq[0].normal_balance()
        if key in ctx.registry and ctx.registry_src.get(key) != self.ws.title:
            return ctx.registry[key]
        if key == "net income" and ctx.registry.get("net income"):
            if ctx.registry_src.get("net income") != self.ws.title:
                return ctx.registry["net income"]
        return None

    def _fuzzy_account(self, label):
        book = self.ctx.book
        k = A.account_key(label)
        if not k or len(k) < 3:
            return None
        if len(k.split()) > 4:
            return None
        hits = [a for key, a in book.accounts.items()
                if key.startswith(k + " ") or k.startswith(key + " ")]
        return hits[0] if len(hits) == 1 else None

    # ------------------------------------------------------------------
    def item_col(self, row):
        if row.nums:
            return max(row.nums)
        if row.holes or row.refill:
            return (row.holes or row.refill)[0]
        # same column as the nearest item with a value (look up, then down)
        rs = sorted(self.rows)
        idx = rs.index(row.r)
        for seq in (reversed(rs[:idx]), rs[idx + 1:]):
            for r in seq:
                other = self.rows[r]
                if other.kind == "wall":
                    break
                if other.kind == "item" and other.nums:
                    return min(other.nums)
        if self.amount_cols:
            return self.amount_cols[0]
        return self._default_col(row)

    def _default_col(self, row):
        ws = self.ws
        c = row.label_col
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= row.r <= rng.max_row and rng.min_col <= c <= rng.max_col:
                c = rng.max_col
        return c + 1

    def components(self, row):
        """Rows that add up to this total, scanning upward."""
        out = []
        side = _bs_side(row.label)
        rs = sorted(r for r in self.rows if r < row.r)
        i = len(rs) - 1
        gap_rows = 0
        last_r = row.r
        cover = row.r
        while i >= 0:
            r = rs[i]
            other = self.rows[r]
            if last_r - r > 2 and out:
                break               # a gap of 2+ empty rows ends the section
            if other.kind == "wall":
                break
            if other.kind == "header":
                cover = r
                break
            if other.kind == "total":
                oside = _bs_side(other.label)
                if side and oside and side != oside and "and" not in norm(other.label):
                    break
                if not _fits_under(self.key(other), self.key(row)):
                    break
                out.append(other)
                cover = other.cover or r
                last_r = cover
                idx = [j for j, rr in enumerate(rs) if rr < cover]
                i = idx[-1] if idx else -1
                continue
            out.append(other)
            cover = r
            last_r = r
            i -= 1
        row.cover = cover
        return list(reversed(out))

    def sign(self, comp, total):
        tkey = self.key(total) if total.r else canonical(total.label)
        clab = norm(comp.label)
        if re.match(r"^(less|deduct|minus)\b", clab):
            neg = True
        elif re.match(r"^(add|plus)\b", clab):
            neg = False
        elif tkey in ("net income", "gross profit", "operating income", "income before taxes"):
            ckey = canonical(comp.label)
            ctype = self.ctx.classify(comp.label)
            if ckey in ("net sales", "total revenues", "gross profit", "operating income",
                        "income before taxes", "net income") or re.search(
                    r"^(total )?(revenues?|sales|income)\b(?!.*expense)", clab) and \
                    "tax" not in clab.split("income", 1)[0]:
                return 1
            neg = ckey in ("cost of goods sold", "total expenses", "total operating expenses") \
                or ctype in (A.EXPENSE, A.COGS, A.CONTRA_REVENUE) or \
                bool(re.search(r"expense|cost of|tax|loss", clab)) and "gross" not in clab \
                and "income" not in clab
        elif tkey == "net sales":
            neg = bool(re.search(r"return|discount|allowance", clab))
        elif tkey in ("ending retained earnings", "ending capital"):
            neg = bool(re.search(r"dividend|drawing|withdraw|net loss|distribution", clab))
        elif tkey == "cost of goods sold":
            neg = bool(re.search(r"ending inventory|inventory, end|return|discount|allowance",
                                 clab))
        elif tkey == "working capital":
            neg = "liabilit" in clab
        else:
            ctype = self.ctx.classify(comp.label)
            neg = ctype in (A.CONTRA_ASSET, A.CONTRA_REVENUE, A.CONTRA_EQUITY, A.CONTRA_COGS)
        if neg and comp.value_col and comp.nums:
            n = self.ctx.number(self.ws, comp.r, comp.value_col)
            if n is not None and n < 0:
                neg = False     # already entered as a negative number
        return -1 if neg else 1

    # ------------------------------------------------------------------
    def run(self):
        self.scan()
        if not self.rows or not any(row.nums or row.holes or row.label for row in
                                    self.rows.values()):
            return
        self.classify_rows()
        for r in sorted(self.rows):
            row = self.rows[r]
            if row.kind == "item":
                self.do_item(row)
            elif row.kind == "total":
                self.do_total(row)

    def _fill(self, r, c, formula, why, fmt=None):
        self.ctx.fill(self.ws, r, c, formula, why, fmt)
        self.ctx.template_fills.add((self.ws.title, f"{get_column_letter(c)}{r}"))

    def _register(self, row, c):
        ref = cell_ref(self.ws, row.r, c)
        key = self.key(row)
        self.ctx.register(key, ref, self.ws)
        if key == "total operating expenses":
            self.ctx.register("total expenses", ref, self.ws)
        if key == "total revenues":
            self.ctx.register("net sales", ref, self.ws)

    def do_item(self, row):
        if row.nums and not row.holes:
            row.value_col = max(row.nums)
            self._register(row, row.value_col)
            return
        if row.nums and row.holes:          # "Less: accum. dep.  (5,000)   ?"  -> subtotal
            row.value_col = max(row.nums)
            self._register(row, row.value_col)
            hole = row.holes[-1]
            if hole > row.value_col:
                self._inline_subtotal(row, hole)
            return
        ratio = ratios.ratio_expr(self.ctx, row.label)
        col = self.item_col(row)
        row.value_col = col
        if ratio:
            expr, fmt, why = ratio
            self._fill(row.r, col, "=" + expr, why, fmt)
            return
        expr = self.link(row)
        if expr:
            self._fill(row.r, col, "=" + expr, f"{row.label} = {_source(expr, self.ws.title)}")
            self._register(row, col)
        elif self.final and (row.holes or row.refill):
            self.ctx.warn(f"{self.ws.title}!{get_column_letter(col)}{row.r}: no data found for "
                          f"'{row.label}' - left blank.")

    def _inline_subtotal(self, row, hole):
        refs = []
        col = row.value_col
        for r in sorted((r for r in self.rows if r <= row.r), reverse=True):
            other = self.rows[r]
            if other.kind in ("header", "wall", "total"):
                break
            if col in other.nums or other.value_col == col:
                s = self.sign(other, Row(0, "total", 0))
                refs.append(("-" if s < 0 else "+") + f"{get_column_letter(col)}{r}")
        if refs:
            f = "=" + "".join(reversed(refs)).lstrip("+")
            self._fill(row.r, hole, f, f"Subtotal of {row.label} section")

    def do_total(self, row):
        comps = self.components(row)
        if row.nums and not row.holes:
            row.value_col = max(row.nums)
            self._register(row, row.value_col)
            return
        if not comps:
            # A total with nothing above it (e.g. on a ratio sheet): link it.
            expr = self.link(row)
            col = (row.holes or row.refill or [self.item_col(row)])[0]
            row.value_col = col
            if expr:
                self._fill(row.r, col, "=" + expr, f"{row.label}: linked")
                self._register(row, col)
            return
        cols = self._target_cols(row, comps)
        for col in cols:
            terms, words = [], []
            for comp in comps:
                cc = self._comp_col(comp, col)
                if cc is None:
                    continue
                s = self.sign(comp, row)
                terms.append(("-" if s < 0 else "+") + f"{get_column_letter(cc)}{comp.r}")
                name = re.sub(r"^(add|less|plus|deduct|minus)\s*:?\s*", "", comp.label.strip(),
                              flags=re.I)
                words.append(("- " if s < 0 else "+ ") + name)
            if not terms:
                continue
            formula = "=" + "".join(terms).lstrip("+")
            names = " ".join(words[:8]) + (" ..." if len(words) > 8 else "")
            if names.startswith("+ "):
                names = names[2:]
            self._fill(row.r, col, formula, f"{row.label} = {names}")
            if row.value_col is None:
                row.value_col = col
                self._register(row, col)

    def _comp_col(self, comp, col):
        if self.comparative:
            return col          # every year column is its own statement
        if comp.value_col is not None:
            return comp.value_col
        if comp.nums:
            return max(c for c in comp.nums if c <= col) if any(c <= col for c in comp.nums) \
                else None
        return None

    def _target_cols(self, row, comps):
        if row.holes or row.refill:
            return row.holes or row.refill
        if self.comparative:
            return [c for c in self.amount_cols if any(c in comp.nums for comp in comps)] or \
                self.amount_cols
        comp_cols = [c.value_col for c in comps if c.value_col]
        if not comp_cols:
            return [self.item_col(row)]
        cc = max(comp_cols)
        items_only = all(c.kind == "item" for c in comps)
        right = [c for c in self.amount_cols if c > cc]
        totals_cols = Counter(r.value_col for r in self.rows.values()
                              if r.kind == "total" and r.value_col and r is not row)
        if items_only and right and (not totals_cols or max(totals_cols) > cc):
            return [right[0]]
        return [cc]


def fill_templates(ctx, final):
    for ws in ctx.wb.worksheets:
        if ws.title in ctx.created or ws.title == "Solver Notes":
            continue
        SheetFiller(ctx, ws, final).run()
